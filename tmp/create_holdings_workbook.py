from __future__ import annotations

from copy import copy
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


BASE_DIR = Path("/Users/yimin.fu/PythonProjects/ctp/docs/holdings")
TODAY = date.today().isoformat()
TARGET = BASE_DIR / f"holdings-{TODAY}.xlsx"

SHEETS = {
    "持仓录入": ["记录ID", "合约代码", "品种名称", "方向", "手数", "开仓均价", "开仓日期", "备注"],
    "原始推荐": [
        "记录ID",
        "推荐日期",
        "推荐时间",
        "合约代码",
        "品种名称",
        "方向",
        "计划类型",
        "原始入场价",
        "原始第一止损位",
        "原始止损位",
        "原始第一止盈位",
        "原始止盈位",
        "原始第一止盈RR",
        "原始准入RR",
        "备注",
    ],
}

WIDTHS = {
    "持仓录入": [12, 14, 14, 10, 10, 12, 12, 24],
    "原始推荐": [12, 12, 10, 14, 14, 10, 12, 12, 14, 12, 14, 12, 12, 12, 24],
}


def latest_previous_workbook() -> Path | None:
    candidates = sorted(BASE_DIR.glob("holdings-*.xlsx"))
    earlier = [path for path in candidates if path.name < TARGET.name]
    return earlier[-1] if earlier else None


def is_non_empty_row(values: list[object]) -> bool:
    for value in values:
        if value is None:
            continue
        if isinstance(value, str) and value.strip() == "":
            continue
        return True
    return False


def header_index_map(ws) -> dict[str, int]:
    mapping = {}
    for idx in range(1, ws.max_column + 1):
        value = ws.cell(1, idx).value
        if isinstance(value, str) and value.strip():
            mapping[value.strip()] = idx
    return mapping


def copy_rows(src_ws, dst_ws, headers: list[str]) -> None:
    if src_ws is None:
        return
    src_map = header_index_map(src_ws)
    for row_idx in range(2, src_ws.max_row + 1):
        values = [src_ws.cell(row_idx, src_map[h]).value if h in src_map else None for h in headers]
        if is_non_empty_row(values):
            dst_ws.append(values)


def style_sheet(ws, headers: list[str], widths: list[int]) -> None:
    ws.freeze_panes = "A2"
    header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    header_font = Font(bold=True)
    for col_idx, (header, width) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(1, col_idx)
        cell.value = header
        cell.fill = copy(header_fill)
        cell.font = copy(header_font)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[cell.column_letter].width = width


def add_dropdowns(ws, column_letter: str, options: list[str]) -> None:
    formula = '"' + ",".join(options) + '"'
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{column_letter}2:{column_letter}1048576")


def main() -> None:
    BASE_DIR.mkdir(parents=True, exist_ok=True)
    if TARGET.exists():
        print(f"exists:{TARGET}")
        return

    source_path = latest_previous_workbook()
    source_wb = load_workbook(source_path) if source_path else None

    wb = Workbook()
    default_ws = wb.active
    wb.remove(default_ws)

    for sheet_name, headers in SHEETS.items():
        ws = wb.create_sheet(sheet_name)
        style_sheet(ws, headers, WIDTHS[sheet_name])
        src_ws = source_wb[sheet_name] if source_wb and sheet_name in source_wb.sheetnames else None
        copy_rows(src_ws, ws, headers)

    add_dropdowns(wb["持仓录入"], "D", ["做多", "做空"])
    add_dropdowns(wb["原始推荐"], "F", ["做多", "做空"])
    add_dropdowns(wb["原始推荐"], "G", ["趋势", "反转"])

    wb.save(TARGET)
    print(f"created:{TARGET}")
    if source_path:
        print(f"source:{source_path}")


if __name__ == "__main__":
    main()
